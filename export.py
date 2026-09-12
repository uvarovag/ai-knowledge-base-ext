"""Rebuild the Excel view from the living knowledge base.

Besides the base columns the sheet carries one column per entry in
config.SOURCE_EXTRA_COLUMNS, so a reviewer can trace an entry back to the
tickets it was built from without opening the JSON.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import common
import config
from common import logger


def format_source_rows(entry: dict[str, Any]) -> str:
    """Render the source row numbers of an entry as a comma separated string."""
    return config.SOURCE_VALUE_SEPARATOR.join(
        str(row) for row in entry.get("source_rows", [])
    )


def build_headers() -> tuple[str, ...]:
    """Return the sheet headers: the base ones plus the extra source columns."""
    return (*config.HEADERS, *config.SOURCE_EXTRA_COLUMNS)


def build_widths() -> tuple[int, ...]:
    """Return one width per header, using a single width for the extra columns."""
    return (
        *config.COLUMN_WIDTHS,
        *(config.EXTRA_COLUMN_WIDTH for _ in config.SOURCE_EXTRA_COLUMNS),
    )


def build_row(entry: dict[str, Any]) -> list[str]:
    """Render one entry as a sheet row, in the order of the headers."""
    source_columns = entry.get("source_columns", {})
    return [
        entry["question"],
        entry["answer"],
        entry["category"],
        entry.get("source_file", ""),
        format_source_rows(entry),
        entry.get("updated_at", ""),
        *(source_columns.get(column, "") for column in config.SOURCE_EXTRA_COLUMNS),
    ]


def write_workbook(entries: list[dict[str, Any]], path: Path) -> None:
    """Write the entries to a formatted single-sheet workbook."""
    headers = build_headers()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = config.SHEET_TITLE

    sheet.append(list(headers))
    header_fill = PatternFill("solid", fgColor=config.HEADER_FILL_COLOR)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for entry in entries:
        sheet.append(build_row(entry))

    for index, width in enumerate(build_widths(), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    wrapped = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrapped

    sheet.freeze_panes = "A2"
    last_column = get_column_letter(len(headers))
    sheet.auto_filter.ref = f"A1:{last_column}{sheet.max_row}"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def run() -> None:
    """Export the living base to the workbook configured in config."""
    entries = common.load_json(config.KNOWLEDGE_BASE_JSON)
    if not entries:
        logger.warning("Knowledge base is empty, nothing to export")
        return

    entries.sort(key=lambda entry: (entry["category"], entry["question"]))
    write_workbook(entries, config.KNOWLEDGE_BASE_XLSX)
    logger.info("Exported %d entries to %s", len(entries), config.KNOWLEDGE_BASE_XLSX)


if __name__ == "__main__":
    common.configure_logging()
    run()
