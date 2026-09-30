"""Mark the questions of a dump whose point cannot be made out, for support.

The pipeline rejects such tickets or turns them into useless entries, and
only the people who answered them can say what was asked. This script asks
the model about every question of every dump in config.DUMP_PATHS and fills
the question cells it finds vague with config.VAGUE_QUESTION_FILL_COLOR, in
the dump itself: the file goes to support as is, and they rewrite the
highlighted questions in place. A cell that was marked on an earlier run and
is clear now loses the fill, so the script can be rerun after their edits.

The model judges the question alone, without the answer: an answer that
explains everything would make a bare topic look like a question, while a
reader of the knowledge base only has the question to go on.

The dump is copied to data/backups/ before it is overwritten. Note that
editing a dump changes its hash, so the pipeline processes it as a new one.
"""

from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from pydantic import BaseModel, Field

import common
import config
import extract
from common import logger

# ----- Prompt --------------------------------------------------------------

SYSTEM_PROMPT_TEMPLATE = """Ты проверяешь вопросы, которые пользователи задали службе поддержки <<DOMAIN>>.
Реши, понятна ли СУТЬ вопроса: что пользователь хочет сделать или узнать.

Смотри только на сам вопрос. Ответа поддержки у тебя нет, и додумывать его
нельзя: вопрос должен быть понятен человеку, который видит только его.

ПОНЯТЕН (question_is_clear = true), если из текста ясно намерение:
- есть вопрос «как», «где», «кто», «можно ли», «что делать»;
- есть просьба или желание: «хочу новый стул», «нужна вода для кулера» —
  пользователь хочет это получить, суть ясна;
- названо, что нужно узнать: «ФИО сервис менеджера» — нужно узнать ФИО.
Небрежность, опечатки, отсутствие знака вопроса, номера заявок и длинные
подробности НЕ делают вопрос непонятным.

НЕПОНЯТЕН (question_is_clear = false), если:
- это только тема или термин без намерения: «Складские запасы», «Номер SKU» —
  неизвестно, посмотреть их, заказать, списать или узнать, что это;
- непонятно, о чём речь: «как заказать?» без предмета;
- вопрос оборван или пуст: «По какой статье закупить/оплатить -».

Если сомневаешься — ставь true: помечать стоит только те вопросы, суть
которых действительно нельзя понять.

ФОРМАТ ОТВЕТА — вызов функции с такими аргументами:
{"question_is_clear": true, "reason": "до 10 слов"}

ПРИМЕР 1
Вопрос: Ремонт в здании
{"question_is_clear": false, "reason": "только тема: ремонт нужен, сроки или к кому идти"}

ПРИМЕР 2
Вопрос: Категория закупки
{"question_is_clear": false, "reason": "только тема: определить, найти или кто отвечает"}

ПРИМЕР 3
Вопрос: Комиссия по списанию ОС
{"question_is_clear": false, "reason": "только тема: кто входит, создать или изменить"}

ПРИМЕР 4
Вопрос: вывоз имущества из ВСП
{"question_is_clear": false, "reason": "только тема без вопроса"}

ПРИМЕР 5
Вопрос: как заказать ? Если не нашел в других инструкциях
{"question_is_clear": false, "reason": "не сказано, что заказать"}

ПРИМЕР 6
Вопрос: По какой статье закупить/оплатить -
{"question_is_clear": false, "reason": "вопрос оборван, предмета нет"}

ПРИМЕР 7
Вопрос: Хочу новый стул
{"question_is_clear": true, "reason": "хочет получить стул"}

ПРИМЕР 8
Вопрос: Нужна вода для куллера
{"question_is_clear": true, "reason": "хочет заказать воду"}

ПРИМЕР 9
Вопрос: ФИО сервис менеджера
{"question_is_clear": true, "reason": "хочет узнать ФИО сервис-менеджера"}

ПРИМЕР 10
Вопрос: По какой статье закупить/оплатить Приобретение POS-терминалов
{"question_is_clear": true, "reason": "нужна статья для покупки терминалов"}

ПРИМЕР 11
Вопрос: Добрый день! По ЗП № 8310000735 подтянулась неправильная форма проекта Договора. Мне нужно в закупай её прикладывать?
{"question_is_clear": true, "reason": "как заменить форму договора"}
"""

USER_PROMPT = """Вопрос: {question}
"""

SYSTEM_PROMPT = common.render_prompt(SYSTEM_PROMPT_TEMPLATE, domain=config.DOMAIN_NAME)


class QuestionVerdict(BaseModel):
    """Решение, понятна ли суть вопроса пользователя."""

    question_is_clear: bool = Field(
        description="Понятно, что пользователь хочет сделать или узнать"
    )
    reason: str = Field(description="Причина решения, до 10 слов")


# ----- Verdicts ------------------------------------------------------------


def load_questions(path: Path) -> dict[int, str]:
    """Return sheet row number -> question text for every row with a question.

    Unlike extract.load_pairs, rows without an answer are kept: support
    rewrites every vague question, answered or not.
    """
    dataframe = pd.read_excel(path, sheet_name=0)
    for column in config.QUESTION_COLUMNS:
        if column not in dataframe.columns:
            raise ValueError(f"Column '{column}' not found in {path.name}")

    questions: dict[int, str] = {}
    for index, row in dataframe.iterrows():
        question = extract.join_columns(row, config.QUESTION_COLUMNS)
        if question:
            questions[int(index) + config.FIRST_DATA_ROW] = question
    logger.info("%s: %d questions", path.name, len(questions))
    return questions


def judge_questions(questions: dict[int, str]) -> dict[int, bool]:
    """Return row number -> whether its question is clear; failed rows are left out.

    Verdicts are cached by question text, so identical questions are asked
    about once and a rerun only pays for the questions support changed. The
    prompt is part of the key: editing its rules or examples invalidates the
    verdicts it produced instead of silently reusing them.
    """
    if config.FORCE_REPROCESS and config.VAGUE_QUESTIONS_CACHE.exists():
        config.VAGUE_QUESTIONS_CACHE.unlink()
    cache = {
        record["key"]: record["clear"]
        for record in common.load_json(config.VAGUE_QUESTIONS_CACHE)
    }

    def save_cache() -> None:
        common.save_json(
            config.VAGUE_QUESTIONS_CACHE,
            [{"key": key, "clear": clear} for key, clear in sorted(cache.items())],
        )

    keys = {
        row: common.hash_text(SYSTEM_PROMPT + "\n" + question)
        for row, question in questions.items()
    }
    pending = {
        keys[row]: (row, question)
        for row, question in questions.items()
        if keys[row] not in cache
    }
    cached_count = sum(key in cache for key in keys.values())
    if cached_count:
        logger.info("%d of %d questions taken from cache", cached_count, len(questions))

    if pending:
        llm = common.build_llm()
        completed = 0
        bar = common.ProgressBar(len(pending), "Vague questions")

        def judge(row: int, question: str) -> dict[str, Any] | None:
            return common.invoke_structured(
                llm,
                SYSTEM_PROMPT,
                USER_PROMPT.format(question=question),
                schema=QuestionVerdict,
                label=f"row {row} question",
            )

        # Results are folded in on the main thread, so the cache needs no lock.
        with ThreadPoolExecutor(max_workers=config.WORKER_COUNT) as executor:
            futures = {
                executor.submit(judge, *item): key for key, item in pending.items()
            }
            for future in as_completed(futures):
                key = futures[future]
                row, question = pending[key]
                verdict = future.result()
                if verdict is None:
                    bar.advance(failed=1)
                    continue
                cache[key] = verdict["question_is_clear"]
                if cache[key]:
                    bar.advance(clear=1)
                else:
                    logger.info(
                        "Row %d vague (%s): %r",
                        row,
                        verdict["reason"],
                        question[:80],
                    )
                    bar.advance(vague=1)
                completed += 1
                if completed % config.SAVE_EVERY == 0:
                    save_cache()
        bar.finish()
        save_cache()

    return {row: cache[key] for row, key in keys.items() if key in cache}


# ----- Marking -------------------------------------------------------------


def mark_dump(path: Path) -> None:
    """Judge every question of one dump and fill the vague ones in place."""
    questions = load_questions(path)
    verdicts = judge_questions(questions)

    workbook = load_workbook(path)
    # load_questions reads the first sheet, so its row numbers refer to it.
    sheet = workbook.worksheets[0]
    headers = {cell.value: cell.column for cell in sheet[1]}
    columns = [headers[name] for name in config.QUESTION_COLUMNS]

    vague_fill = PatternFill("solid", fgColor=config.VAGUE_QUESTION_FILL_COLOR)
    marked = 0
    for row, is_clear in verdicts.items():
        for column in columns:
            cell = sheet.cell(row=row, column=column)
            if not is_clear:
                cell.fill = vague_fill
            elif cell.fill.fgColor.rgb == vague_fill.fgColor.rgb:
                # Marked on an earlier run and rewritten since.
                cell.fill = PatternFill()
        marked += not is_clear

    backup_path = common.backup_file(path)
    logger.info("Dump backed up to %s", backup_path)
    workbook.save(path)
    logger.info(
        "%s: %d vague questions marked, %d clear, %d not judged",
        path.name,
        marked,
        len(verdicts) - marked,
        len(questions) - len(verdicts),
    )


if __name__ == "__main__":
    common.configure_logging()
    for dump_path in config.DUMP_PATHS:
        mark_dump(dump_path)
