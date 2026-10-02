"""Excel input and output: reading question/answer pairs, writing views of a base."""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
from kb.utils.logs import logger

# ----- Reading -------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class SourcePair:
    """One question/answer row of a source sheet; either side may be empty."""

    row_number: int
    question: str
    answer: str
    source_columns: dict[str, str] = field(default_factory=dict)


def join_columns(row: pd.Series, columns: Sequence[str]) -> str:
    """Join the non-empty values of the given columns into a single text."""
    parts = [
        str(row[column]).strip()
        for column in columns
        if pd.notna(row[column]) and str(row[column]).strip()
    ]
    return config.COLUMN_SEPARATOR.join(parts)


def read_pairs(
    path: Path,
    question_columns: Sequence[str],
    answer_columns: Sequence[str],
    extra_columns: Sequence[str],
) -> list[SourcePair]:
    """Read the first sheet and return every row with a question or an answer.

    A missing question or answer column is an error. A missing extra column
    only costs traceability, so it is skipped with a note: an older file may
    lack it.
    """
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    logger.info("Reading %s", path)
    dataframe = pd.read_excel(path, sheet_name=0)
    logger.info("Loaded %d rows, %d columns", len(dataframe), len(dataframe.columns))

    for column in (*question_columns, *answer_columns):
        if column not in dataframe.columns:
            raise ValueError(
                f"Column '{column}' not found in {path.name}. "
                f"Available: {dataframe.columns.tolist()}"
            )
    available_extra = [column for column in extra_columns if column in dataframe.columns]
    missing_extra = [column for column in extra_columns if column not in available_extra]
    if missing_extra:
        # INFO: a config mismatch for the person at the terminal, not a failure.
        logger.info("Extra columns not in %s, skipped: %s", path.name, missing_extra)

    pairs: list[SourcePair] = []
    for index, row in dataframe.iterrows():
        question = join_columns(row, question_columns)
        answer = join_columns(row, answer_columns)
        if question or answer:
            pairs.append(
                SourcePair(
                    row_number=int(index) + config.FIRST_DATA_ROW,
                    question=question,
                    answer=answer,
                    source_columns={
                        column: str(row[column]).strip()
                        for column in available_extra
                        if pd.notna(row[column]) and str(row[column]).strip()
                    },
                )
            )
    return pairs


# ----- Writing -------------------------------------------------------------


def write_sheet(
    path: Path,
    headers: Sequence[str],
    widths: Sequence[int],
    rows: list[list[Any]],
    row_fills: Sequence[str | None] = (),
) -> None:
    """Write one formatted sheet: bold filled header, wrapped cells, filter.

    row_fills gives a fill colour per data row, None for none.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = config.SHEET_TITLE

    sheet.append(list(headers))
    header_fill = PatternFill("solid", fgColor=config.HEADER_FILL_COLOR)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        sheet.append(row)

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    wrapped = Alignment(wrap_text=True, vertical="top")
    for sheet_row in sheet.iter_rows(min_row=2):
        for cell in sheet_row:
            cell.alignment = wrapped
    for sheet_row, color in zip(sheet.iter_rows(min_row=2), row_fills):
        if color:
            for cell in sheet_row:
                cell.fill = PatternFill("solid", fgColor=color)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def read_knowledge_base(path: Path) -> list[dict[str, Any]]:
    """Read back a base this project wrote (write_knowledge_base), edits included.

    The sheet holds every field of an entry, so a reviewer's edits — a row
    deleted, a text corrected, a doubt reason cleared — become the entries.
    Columns are found by their headers; every column besides config.HEADERS
    is an extra source column. A row without a question or an answer is
    dropped: the reviewer emptied it.
    """
    question, answer, category, doubtful, source_file, source_rows, updated = (
        config.HEADERS
    )
    dataframe = pd.read_excel(path, sheet_name=0, dtype=str, keep_default_na=False)
    missing = [column for column in (question, answer, category) if column not in dataframe]
    if missing:
        raise SystemExit(
            f"{path.name} is not a base this project wrote: no columns {missing}. "
            "Repair it first with make repair-base."
        )
    extra_columns = [column for column in dataframe.columns if column not in config.HEADERS]

    def cell(row: pd.Series, column: str) -> str:
        return str(row.get(column, "")).strip()

    entries: list[dict[str, Any]] = []
    for _, row in dataframe.iterrows():
        if not cell(row, question) or not cell(row, answer):
            continue
        entry: dict[str, Any] = {
            "category": cell(row, category),
            "question": cell(row, question),
            "answer": cell(row, answer),
            "source_file": cell(row, source_file),
            "source_rows": [
                int(part)
                for part in cell(row, source_rows).split(config.SOURCE_VALUE_SEPARATOR.strip())
                if part.strip().isdigit()
            ],
            "source_columns": {
                column: cell(row, column) for column in extra_columns if cell(row, column)
            },
            "updated_at": cell(row, updated),
        }
        if cell(row, doubtful):
            entry["doubtful"] = cell(row, doubtful)
        entries.append(entry)
    return entries


def write_knowledge_base(entries: list[dict[str, Any]], path: Path) -> None:
    """Write a base as a review-friendly sheet, sorted by category and question.

    Besides the base columns the sheet carries one column per extra source
    column the entries hold, in the order first seen, so a reviewer can trace
    an entry back to its tickets. They are taken from the entries, not from a
    config: a base grown from dumps with different columns keeps all of them.
    A doubtful entry's row is highlighted, its reason in its own column.
    """
    entries = sorted(entries, key=lambda entry: (entry["category"], entry["question"]))
    extra_columns = list(
        dict.fromkeys(
            column for entry in entries for column in entry.get("source_columns", {})
        )
    )
    rows = [
        [
            entry["question"],
            entry["answer"],
            entry["category"],
            entry.get("doubtful", ""),
            entry.get("source_file", ""),
            config.SOURCE_VALUE_SEPARATOR.join(
                str(row) for row in entry.get("source_rows", [])
            ),
            entry.get("updated_at", ""),
            *(entry.get("source_columns", {}).get(column, "") for column in extra_columns),
        ]
        for entry in entries
    ]
    write_sheet(
        path,
        (*config.HEADERS, *extra_columns),
        (*config.COLUMN_WIDTHS, *(config.EXTRA_COLUMN_WIDTH for _ in extra_columns)),
        rows,
        [config.DOUBTFUL_FILL_COLOR if entry.get("doubtful") else None for entry in entries],
    )
    doubtful = sum(bool(entry.get("doubtful")) for entry in entries)
    logger.info(
        "Exported %d entries, %d of them doubtful, to %s", len(entries), doubtful, path
    )


def write_rejected(records: list[dict[str, Any]], path: Path) -> None:
    """Write the rows left out of a base, with the reason, for a reviewer."""
    rows = [
        [
            record["source_row"],
            record["reason"],
            record["question"],
            record["answer"],
            record.get("model_question") or "",
            record.get("model_answer") or "",
        ]
        for record in sorted(records, key=lambda record: record["source_row"])
    ]
    write_sheet(path, config.REJECTED_HEADERS, config.REJECTED_COLUMN_WIDTHS, rows)
    logger.info("Exported %d rejected rows to %s", len(records), path)
